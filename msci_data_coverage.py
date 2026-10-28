#!/usr/bin/env python3
"""Extract the "Closed-end fund count and capitalization by asset class" table from the
Data Coverage page of an MSCI Private Capital Benchmarks Summary PDF into an xlsx file.

The table is found by its header row (Asset class, Funds, FoFs, Total, Cap (USD M)) on the
left half of the page. Each number is matched to the column whose header it lines up
under, and the hierarchy (asset class > sub-class) comes from each label's indent.
Sub-classes are named with their asset class, e.g. "Equity | Buyout".
A "-" in the PDF is written as an empty cell.

The script checks that Funds + FoFs = Total on every row and that each group's
sub-classes add up to the group (allowing for rounding in Cap), and prints a note for
any that don't. The values are written exactly as printed either way: in the Q2 2026
report, for example, Credit's 48 FoFs aren't broken down by sub-class.

Usage:
    python msci_data_coverage.py report.pdf [-o data_coverage.xlsx] [--page 6]
"""

import argparse
import re
import sys

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TITLE = "Closed-end fund count and capitalization by asset class"
VALUE_COLUMNS = ["Funds", "FoFs", "Total", "Cap (USD M)"]
NUMBER = re.compile(r"^(-|\d{1,3}(,\d{3})*)$")
ROW_TOLERANCE = 2.5  # points: words this close vertically are on the same row
ALIGN_TOLERANCE = 4.0  # points: a number's right edge vs. its column header's right edge


def group_rows(words):
    """Group words into rows by their top coordinate, top to bottom."""
    rows = []
    for word in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if rows and abs(word["top"] - rows[-1][0]["top"]) <= ROW_TOLERANCE:
            rows[-1].append(word)
        else:
            rows.append([word])
    return [sorted(row, key=lambda w: w["x0"]) for row in rows]


def find_header(rows, page_width):
    """Return (index of the header row, right edge of each value column, label column end)."""
    for i, row in enumerate(rows):
        text = " ".join(w["text"] for w in row)
        if row[0]["x0"] < page_width / 2 and text.startswith("Asset class Funds FoFs Total Cap"):
            by_text = {w["text"]: w for w in row}
            right_edges = [by_text["Funds"]["x1"], by_text["FoFs"]["x1"],
                           by_text["Total"]["x1"], by_text["M)"]["x1"]]
            return i, right_edges, by_text["Funds"]["x0"]
    sys.exit("Couldn't find the table header (Asset class, Funds, FoFs, Total, Cap (USD M)) "
             "on that page.")


def parse_row(row, right_edges, label_end):
    """Split a row into (label words, {column: value}). Returns None if the row doesn't
    fit the table: numbers that don't line up under a column, or text in a number column."""
    label = [w for w in row if w["x1"] <= label_end]
    values = {}
    for word in row:
        if word["x1"] <= label_end:
            continue
        if not NUMBER.match(word["text"]):
            return None
        column = min(range(len(right_edges)), key=lambda c: abs(word["x1"] - right_edges[c]))
        if abs(word["x1"] - right_edges[column]) > ALIGN_TOLERANCE or column in values:
            return None
        values[column] = None if word["text"] == "-" else int(word["text"].replace(",", ""))
    if not label or len(values) != len(right_edges):
        return None
    return label, values


def extract_table(pdf_path, page_number):
    with pdfplumber.open(pdf_path) as pdf:
        if not 1 <= page_number <= len(pdf.pages):
            sys.exit(f"The PDF has {len(pdf.pages)} pages; there is no page {page_number}.")
        page = pdf.pages[page_number - 1]
        # Only the left half: the page has a second table on the right.
        words = [w for w in page.extract_words(extra_attrs=["fontname"])
                 if w["x0"] < page.width / 2]
        width = page.width

    rows = group_rows(words)
    header, right_edges, label_end = find_header(rows, width)
    # Keep only words within the table's own width (from its header to the right edge
    # of the Cap column), so nothing from the table on the right gets mixed in.
    left, right = rows[header][0]["x0"] - 5, right_edges[-1] + ALIGN_TOLERANCE
    rows = group_rows([w for row in rows[header + 1:] for w in row
                       if w["x0"] >= left and w["x1"] <= right])

    records = []
    for row in rows:
        parsed = parse_row(row, right_edges, label_end)
        if parsed is None:
            break  # past the end of the table (the vintage chart starts below it)
        label, values = parsed
        records.append({
            "label": " ".join(w["text"] for w in label),
            "indent": label[0]["x0"],
            "bold": any("Bold" in w["fontname"] for w in label),
            "values": [values[c] for c in range(len(right_edges))],
        })
    if not records:
        sys.exit("Found the header but no table rows under it.")

    # Indent positions, smallest first, give the hierarchy levels.
    indents = []
    for x in sorted(r["indent"] for r in records):
        if not indents or x - indents[-1] > 3:
            indents.append(x)
    for r in records:
        r["level"] = max(i for i, x in enumerate(indents) if r["indent"] >= x - 3)

    # Parent = nearest row above with a smaller level.
    stack = []
    for r in records:
        while stack and stack[-1]["level"] >= r["level"]:
            stack.pop()
        r["parent"] = stack[-1]["label"] if stack else ""
        stack.append(r)

    # Sub-classes (level 2) are named with their asset class: "Equity | Buyout".
    for r in records:
        if r["level"] == 2:
            r["label"] = f"{r['parent']} | {r['label']}"
    return records


def check_totals(records):
    """Note rows where Funds + FoFs != Total, or where a group's rows don't add up.
    Cap is in rounded millions, so it may be off by up to one per row added."""
    problems = 0
    for r in records:
        funds, fofs, total, _ = (v or 0 for v in r["values"])
        if funds + fofs != total:
            print(f"Note: in the PDF, {r['label']}: Funds {funds:,} + FoFs {fofs:,} "
                  f"!= Total {total:,}", file=sys.stderr)
            problems += 1
    for i, r in enumerate(records):
        children = []
        for child in records[i + 1:]:
            if child["level"] <= r["level"]:
                break
            if child["level"] == r["level"] + 1:
                children.append(child)
        if not children:
            continue
        for c, name in enumerate(VALUE_COLUMNS):
            expected = r["values"][c] or 0
            added = sum(child["values"][c] or 0 for child in children)
            rounding = len(children) if name == "Cap (USD M)" else 0
            if abs(added - expected) > rounding:
                print(f"Note: in the PDF, {r['label']} {name} is {expected:,} but the rows "
                      f"under it add up to {added:,}", file=sys.stderr)
                problems += 1
    return problems


def write_xlsx(records, path, source):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data coverage"

    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = source
    ws["A2"].font = Font(italic=True, color="666666")

    header_row = 4
    headers = ["Asset class", "Parent", "Level"] + VALUE_COLUMNS
    thin = Side(style="thin", color="999999")
    for c, name in enumerate(headers, start=1):
        cell = ws.cell(header_row, c, name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE4F0")
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(horizontal="left" if c <= 2 else "right")

    for i, r in enumerate(records, start=header_row + 1):
        font = Font(bold=r["bold"])
        label = ws.cell(i, 1, r["label"])
        label.alignment = Alignment(indent=2 * r["level"])
        label.font = font
        ws.cell(i, 2, r["parent"]).font = font
        ws.cell(i, 3, r["level"]).font = font
        for c, value in enumerate(r["values"], start=4):
            cell = ws.cell(i, c, value)
            cell.number_format = "#,##0"
            cell.font = font

    for column, width in zip("ABCDEFG", (42, 20, 7, 10, 10, 10, 14)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f"A{header_row}:G{header_row + len(records)}"
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pdf", help="MSCI Private Capital Benchmarks Summary PDF")
    parser.add_argument("-o", "--output", default="data_coverage.xlsx",
                        help="output xlsx path (default: data_coverage.xlsx)")
    parser.add_argument("--page", type=int, default=6,
                        help="page number of the Data Coverage page (default: 6)")
    args = parser.parse_args()

    records = extract_table(args.pdf, args.page)
    problems = check_totals(records)
    write_xlsx(records, args.output, f"Source: {args.pdf}, page {args.page}")
    print(f"Saved {len(records)} rows to {args.output}"
          + (f"; {problems} total(s) in the PDF don't add up (see notes above)" if problems
             else "; all totals add up"))


if __name__ == "__main__":
    main()
